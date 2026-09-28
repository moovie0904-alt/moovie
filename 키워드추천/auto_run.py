# -*- coding: utf-8 -*-
"""
자동 실행: 엑셀에서 C열(핵심키워드)이 빈 줄을 찾아
  1) 내 PC의 Claude Code(claude 명령)에게 핵심키워드를 뽑게 해서 C열에 넣고
  2) 네이버 검색광고 API로 경쟁 낮은 추천키워드 3개를 D열에 넣은 뒤
  3) 원래 엑셀 파일(C:\자동화시스템 의 '핵심키워드' 엑셀)에 그대로 저장합니다. 저장 전 백업 폴더에 복사본을 남깁니다.

사용법:
    python auto_run.py                         (기본 파일, 빈 줄 200개)
    python auto_run.py --count 50              (50개만)
    python auto_run.py "D:\파일.xlsx"
"""
import json
import os
import shutil
import subprocess
import sys
from pathlib import Path

from openpyxl import load_workbook

import recommend_keywords as rk

NAME_COL, TAG_COL = 1, 2          # A열 상품명, B열 태그
CHUNK = 50                        # Claude에게 한 번에 보내는 상품 수

PROMPT = """너는 네이버 스마트스토어 상품 등록용 키워드 전문가다.
입력(표준입력)의 각 줄은 "행번호<TAB>상품명<TAB>태그" 이다.
각 행마다 상품명과 태그를 보고, 구매자가 네이버 쇼핑에서 가장 많이 검색할 '핵심키워드' 1개를 고른다.

기준:
- "무엇을 파는지"를 나타내는 제품 종류 이름. 상품명과 태그에 함께 나오는 말을 우선한다.
- 색상·크기·수량·재질·꾸밈말(모던, 만능 등)·사용장소(원룸, 업소용 등)는 빼되, 제품 종류 구분에 꼭 필요하면 붙인다.
- "건조기", "사다리"처럼 너무 넓은 말보다 "식품건조기", "다락방사다리"처럼 한 단계 구체적으로.
- 브랜드·캐릭터·차종·연예인·게임 이름 같은 상표는 절대 쓰지 않고 일반 명사로 바꾼다 (예: 몰텐 축구공 → 5호축구공, 카니발 보조의자 → 차량용보조의자).
- 띄어쓰기 없이 한국어, 보통 2~10글자.

예시: 키위 플라스틱 트레이 커버포함 36구 / 태그 포장,과일트레이,과일상자 → 과일트레이
예시: 황금 골드 세면볼 수도꼭지 코브라 씽크 주방 / 태그 거위목,싱크대,수전 → 골드주방수전
예시: 가정용 고추건조기 소형 고구마 과일 간식 건조기 / 태그 채소,수제간식,건과일 → 식품건조기

파일을 읽거나 쓰지 말고, 도구를 쓰지 말고, 입력의 모든 행번호에 대해 결과만 돌려준다."""

SCHEMA = {
    "type": "object",
    "properties": {"items": {"type": "array", "items": {
        "type": "object",
        "properties": {"row": {"type": "integer"}, "keyword": {"type": "string"}},
        "required": ["row", "keyword"]}}},
    "required": ["items"],
}


def find_claude() -> str:
    exe = shutil.which("claude")
    if exe:
        return exe
    guess = Path(os.environ.get("USERPROFILE", "")) / ".local" / "bin" / "claude.exe"
    if guess.exists():
        return str(guess)
    sys.exit("claude 명령을 찾을 수 없습니다. Claude Code를 설치했는지 확인하세요 (설치방법.txt 참고).")


def ask_claude(claude: str, lines: list[str]) -> dict[int, str]:
    proc = subprocess.run(
        [claude, "-p", PROMPT, "--output-format", "json", "--json-schema", json.dumps(SCHEMA, ensure_ascii=False)],
        input="\n".join(lines), capture_output=True, text=True, encoding="utf-8", timeout=900,
    )
    try:
        data = json.loads(proc.stdout)
    except json.JSONDecodeError:
        print("  Claude 응답을 읽지 못했습니다:", (proc.stdout or proc.stderr)[:300])
        return {}
    if data.get("is_error"):
        print("  Claude 오류:", str(data.get("result"))[:300])
        return {}
    out = data.get("structured_output")
    if out is None:  # 구조화 출력이 없으면 result 텍스트에서 JSON을 찾아봄
        try:
            out = json.loads(data.get("result", ""))
        except (json.JSONDecodeError, TypeError):
            print("  Claude 응답 형식이 예상과 다릅니다:", str(data.get("result"))[:300])
            return {}
    return {int(i["row"]): "".join(str(i["keyword"]).split()) for i in out.get("items", []) if i.get("keyword")}


def clean(v) -> str:
    return " ".join(str(v or "").split())


def main() -> None:
    args = sys.argv[1:]
    count = 200
    if "--count" in args:
        i = args.index("--count")
        count = int(args[i + 1])
        del args[i:i + 2]
    in_path = Path(args[0]) if args else rk.default_input()
    if not in_path.exists():
        sys.exit(f"파일을 찾을 수 없습니다: {in_path}")

    claude = find_claude()
    print(f"읽는 중: {in_path} (상품이 많으면 1~2분 걸립니다)")
    wb = load_workbook(in_path)
    ws = wb.active
    ws.cell(1, rk.KEYWORD_COL).value = ws.cell(1, rk.KEYWORD_COL).value or "핵심키워드"

    # 1) C열이 빈 줄 찾기
    todo = []
    for r in range(2, ws.max_row + 1):
        if len(todo) >= count:
            break
        if ws.cell(r, NAME_COL).value and not clean(ws.cell(r, rk.KEYWORD_COL).value):
            todo.append(r)
    print(f"핵심키워드를 채울 줄: {len(todo)}개" + (f" ({todo[0]}행 ~ {todo[-1]}행)" if todo else ""))

    filled_rows = []
    try:
        for s in range(0, len(todo), CHUNK):
            rows = todo[s:s + CHUNK]
            lines = [f"{r}\t{clean(ws.cell(r, NAME_COL).value)}\t{clean(ws.cell(r, TAG_COL).value)}" for r in rows]
            print(f"  Claude에게 요청 중... {rows[0]}~{rows[-1]}행")
            got = ask_claude(claude, lines)
            for r in rows:
                if got.get(r):
                    ws.cell(r, rk.KEYWORD_COL).value = got[r]
                    filled_rows.append(r)
            print(f"    {sum(1 for r in rows if got.get(r))}/{len(rows)}개 채움")

        # 2) D열 추천키워드 (naver_key.txt 가 있을 때만)
        if (rk.HERE / "naver_key.txt").exists():
            need = {}
            for r in range(2, ws.max_row + 1):
                core = clean(ws.cell(r, rk.KEYWORD_COL).value)
                if core and not ws.cell(r, rk.RESULT_COL).value:
                    need.setdefault(core, []).append(r)
            rk.fill_recommendations(ws, rk.load_keys(), need)
        else:
            print("※ naver_key.txt 가 없어 D열(추천키워드)은 건너뜁니다.")
    except KeyboardInterrupt:
        print("\n중단됨 - 지금까지 결과를 저장합니다.")
    finally:
        print("저장 중...")
        rk.backup_and_save(wb, in_path)


if __name__ == "__main__":
    main()
