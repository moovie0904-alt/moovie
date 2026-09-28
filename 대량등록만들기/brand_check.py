# -*- coding: utf-8 -*-
"""
결과\키워드*\대량등록*.xlsx 의 A열 상품명 중 브랜드가 의심되는 것을 찾습니다.

  - 판정: ① 브랜드목록.txt 에 적은 단어가 들어 있으면 바로 의심
          ② 나머지는 내 PC의 Claude Code(claude 명령)에게 250개씩 물어봄
  - 표시: 의심 상품의 A열 칸을 빨간색으로 칠함 (상품명 글자는 그대로)
  - 목록: 브랜드의심_목록.xlsx (파일, 행, 상품명, 의심 단어)

한 번 실행에 파일 100개씩 검사합니다. 다시 실행하면 검사 안 한 파일부터 이어서 합니다.
    python brand_check.py              (파일 100개)
    python brand_check.py --files 300  (파일 300개)
"""
import csv
import io
import json
import os
import shutil
import subprocess
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

HERE = Path(__file__).resolve().parent
RESULT_DIR = HERE / "결과"
DONE_FILE = HERE / "브랜드검사_완료.txt"          # 검사 끝난 파일 목록
CACHE_FILE = HERE / "브랜드검사_판정.csv"         # 상품명 → 의심 단어 (빈칸 = 브랜드 아님)
SUSPECT_FILE = HERE / "브랜드검사_의심.csv"       # 파일, 행, 상품명, 의심 단어
REPORT_FILE = HERE / "브랜드의심_목록.xlsx"
MY_BRANDS = HERE / "브랜드목록.txt"
FILES_PER_RUN = 100
CHUNK = 250
ROWS = range(2, 52)
RED = PatternFill("solid", fgColor="FFC7CE")

PROMPT = """너는 네이버 스마트스토어 상표권 검수 담당자다.
표준입력의 각 줄은 "번호<TAB>상품명" 이다.
상품명에 브랜드명·회사명·상표·캐릭터·연예인/선수 이름·차종/모델명·게임/애니/드라마 제목이 들어 있어
상표권 문제가 될 수 있는 상품만 골라, 그 번호와 문제 단어를 돌려준다.
"호환", "용" 처럼 다른 브랜드 제품용이라는 표현이어도 브랜드명이 들어 있으면 포함한다.
일반 단어와 같은 이름이라도 상품 맥락상 브랜드·모델을 가리키면 포함한다.
예: 카니발·싼타페·모닝·레이(차종), 갤럭시·아이폰(기기), 다이슨·샥즈(회사), 포켓몬·쿠로미(캐릭터), 줌바(상표).
일반 명사로만 된 상품명(예: 원형화로테이블 고기불판, 차량용 보조의자)은 넣지 않는다.
파일을 읽거나 쓰지 말고, 도구를 쓰지 말고, 결과만 돌려준다."""
SCHEMA = {"type": "object", "properties": {"items": {"type": "array", "items": {
    "type": "object", "properties": {"id": {"type": "integer"}, "brand": {"type": "string"}},
    "required": ["id", "brand"]}}}, "required": ["items"]}


def find_claude():
    exe = shutil.which("claude")
    if exe:
        return exe
    guess = Path(os.environ.get("USERPROFILE", "")) / ".local" / "bin" / "claude.exe"
    if guess.exists():
        return str(guess)
    sys.exit("claude 명령을 찾을 수 없습니다. Claude Code를 먼저 설치하세요 (키워드추천\\설치방법.txt 참고).")


def load_cache() -> dict[str, str]:
    if not CACHE_FILE.exists():
        return {}
    with open(CACHE_FILE, encoding="utf-8-sig", newline="") as f:
        return {row[0]: row[1] for row in csv.reader(f) if len(row) >= 2}


def append_csv(path: Path, rows) -> None:
    new = not path.exists()
    with open(path, "a", encoding="utf-8-sig" if new else "utf-8", newline="") as f:
        csv.writer(f).writerows(rows)


def my_brands() -> list[str]:
    if not MY_BRANDS.exists():
        return []
    return ["".join(l.split()) for l in MY_BRANDS.read_text(encoding="utf-8-sig").splitlines()
            if l.strip() and not l.startswith("#")]


def ask_claude(claude: str, names: list[str]) -> dict[str, str]:
    lines = [f"{i}\t{n}" for i, n in enumerate(names)]
    try:
        proc = subprocess.run([claude, "-p", PROMPT, "--output-format", "json", "--json-schema", json.dumps(SCHEMA)],
                              input="\n".join(lines), capture_output=True, text=True, encoding="utf-8", timeout=900)
        data = json.loads(proc.stdout)
        if data.get("is_error"):
            print("  Claude 오류:", str(data.get("result"))[:300])
            return None
        out = data.get("structured_output") or json.loads(data.get("result", "{}"))
    except (json.JSONDecodeError, subprocess.TimeoutExpired, TypeError, AttributeError) as e:
        print(f"  Claude 응답 문제 ({e}) - 이 묶음은 다음 실행 때 다시 합니다.")
        return None
    result = {n: "" for n in names}
    for it in out.get("items", []):
        i = int(it.get("id", -1))
        if 0 <= i < len(names) and it.get("brand"):
            result[names[i]] = str(it["brand"]).strip()
    return result


def mark_file(path: Path, rows: list[int]) -> None:
    wb = load_workbook(path)
    ws = wb.worksheets[0]
    for r in rows:
        ws[f"A{r}"].fill = RED
    tmp = path.with_name(path.stem + ".tmp.xlsx")
    wb.save(tmp)
    os.replace(tmp, path)


def write_report() -> None:
    if not SUSPECT_FILE.exists():
        return
    wb = Workbook()
    ws = wb.active
    ws.title = "브랜드의심"
    ws.append(["파일", "행", "상품명", "의심 단어"])
    with open(SUSPECT_FILE, encoding="utf-8-sig", newline="") as f:
        for row in csv.reader(f):
            if len(row) >= 4:
                ws.append([row[0], int(row[1]), row[2], row[3]])
    for col, w in zip("ABCD", (34, 6, 60, 20)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    try:
        wb.save(REPORT_FILE)
    except PermissionError:
        print(f"※ {REPORT_FILE.name} 이 열려 있어 목록을 갱신하지 못했습니다. 닫고 다시 실행하세요.")


def main() -> None:
    args = sys.argv[1:]
    per_run = int(args[args.index("--files") + 1]) if "--files" in args else FILES_PER_RUN
    claude = find_claude()
    files = sorted(p for p in RESULT_DIR.glob("*/대량등록*.xlsx") if not p.name.endswith(".tmp.xlsx"))
    done = set(DONE_FILE.read_text(encoding="utf-8").splitlines()) if DONE_FILE.exists() else set()
    todo = [p for p in files if str(p.relative_to(HERE)) not in done][:per_run]
    print(f"전체 파일 {len(files):,}개 / 검사 완료 {len(done):,}개 / 이번에 검사 {len(todo):,}개")
    if not todo:
        write_report()
        print("검사할 파일이 없습니다.")
        return

    # 1) 이번 파일들의 상품명 모으기
    names_by_file = {}
    for p in todo:
        ws = load_workbook(p, read_only=True).worksheets[0]
        names_by_file[p] = [(r, " ".join(str(v).split())) for r, (v,) in
                            enumerate(ws.iter_rows(min_row=2, max_row=51, max_col=1, values_only=True), 2) if v]

    # 2) 판정 (브랜드목록 → 캐시 → Claude)
    cache, brands = load_cache(), my_brands()
    unknown = []
    for rows in names_by_file.values():
        for _, n in rows:
            if n in cache or n in unknown:
                continue
            hit = next((b for b in brands if b and b in "".join(n.split())), None)
            if hit:
                cache[n] = hit
                append_csv(CACHE_FILE, [[n, hit]])
            else:
                unknown.append(n)
    print(f"Claude에게 물어볼 상품명 {len(unknown):,}개")
    for s in range(0, len(unknown), CHUNK):
        chunk = unknown[s:s + CHUNK]
        print(f"  판정 중... {s + len(chunk):,}/{len(unknown):,}")
        got = ask_claude(claude, chunk)
        if got is None:
            continue
        cache.update(got)
        append_csv(CACHE_FILE, [[n, b] for n, b in got.items()])

    # 3) 파일 표시 + 목록 (판정이 다 끝난 파일만 완료 처리)
    found = 0
    for p, rows in names_by_file.items():
        if any(n not in cache for _, n in rows):
            continue                                   # 판정 못 받은 상품이 있으면 다음에 다시
        sus = [(r, n, cache[n]) for r, n in rows if cache[n]]
        if sus:
            mark_file(p, [r for r, _, _ in sus])
            append_csv(SUSPECT_FILE, [[str(p.relative_to(RESULT_DIR)), r, n, b] for r, n, b in sus])
            found += len(sus)
        with open(DONE_FILE, "a", encoding="utf-8") as f:
            f.write(str(p.relative_to(HERE)) + "\n")
    write_report()
    print(f"\n이번 실행: 브랜드 의심 {found:,}개 표시 (빨간 칸). 목록: {REPORT_FILE.name}")
    left = len(files) - len(done) - len(todo)
    if left > 0:
        print(f"남은 파일 {left:,}개 - 다시 실행하면 이어서 검사합니다.")


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        print("\n중단됨 - 다시 실행하면 이어서 검사합니다.")
