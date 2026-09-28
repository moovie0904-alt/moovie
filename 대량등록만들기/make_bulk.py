# -*- coding: utf-8 -*-
"""
키워드A.xlsx / 키워드B.xlsx 의 상품을 49개씩 나눠
대량등록 양식(대량등록001.xlsx)에 붙여넣은 파일을 만듭니다.

  원본 A열 상품명  → 양식 A열 (상품명)
  원본 B열 수집URL → 양식 B열 (수집 URL)
  원본 C열 배송비  → 양식 D열 (배송비)
  양식의 머리글·안내문(G열)·테두리 등은 그대로 유지

결과: 결과\키워드A\대량등록0001.xlsx, 대량등록0002.xlsx ...
      결과\키워드B\대량등록0001.xlsx ...

중간에 멈춰도 다시 실행하면 이미 만든 파일은 건너뛰고 이어서 만듭니다.
(파일은 임시 이름으로 저장한 뒤 이름을 바꾸므로, 만들다 만 파일은 남지 않습니다)
"""
import io
import os
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

# ───────── 설정 ─────────
PER_FILE = 49                    # 파일 하나에 넣을 상품 수 (양식 최대 50개)
FIRST_ROW = 2                    # 양식에서 상품을 넣기 시작하는 행
CLEAR_ROWS = range(2, 52)        # 양식의 예시 내용을 지울 행 (2~51행)
CLEAR_COLS = "ABCDEF"            # 예시 내용을 지울 열 (G열 안내문은 유지)
OUT_PREFIX = "대량등록"
# ────────────────────────

HERE = Path(__file__).resolve().parent
OUT_DIR = HERE / "결과"


def find_files():
    templates = [p for p in HERE.glob("대량등록*.xlsx") if not p.name.startswith("~$")]
    if len(templates) != 1:
        sys.exit("이 폴더에 대량등록 양식 파일(예: 대량등록001.xlsx)이 1개 있어야 합니다. 지금: "
                 + (", ".join(p.name for p in templates) or "없음"))
    sources = sorted(p for p in HERE.glob("키워드*.xlsx") if not p.name.startswith("~$"))
    if not sources:
        sys.exit("이 폴더에 키워드A.xlsx, 키워드B.xlsx 같은 원본 파일이 없습니다.")
    return templates[0], sources


def find_cols(header) -> tuple[int, int, int]:
    """머리글 이름으로 상품명/URL/배송비 열 위치를 찾음. 못 찾으면 A/B/C열."""
    names = ["".join(str(h or "").split()) for h in header]

    def idx(keys, default):
        for i, n in enumerate(names):
            if any(k in n for k in keys):
                return i
        return default
    return idx(["상품명"], 0), idx(["URL", "url", "주소"], 1), idx(["배송비"], 2)


def read_products(path: Path):
    """원본을 한 줄씩 읽어 (상품명, URL, 배송비) 를 돌려줌. URL이 빈 줄은 건너뜀."""
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows = ws.iter_rows(values_only=True)
    header = next(rows, ())
    ci_name, ci_url, ci_ship = find_cols(header)
    for row in rows:
        def get(i):
            return row[i] if i < len(row) else None
        url = str(get(ci_url) or "").strip()
        if not url:
            continue
        name = " ".join(str(get(ci_name) or "").split())
        ship = get(ci_ship)
        if isinstance(ship, float) and ship.is_integer():
            ship = int(ship)
        ship = "" if ship is None else str(ship).strip()
        yield name, url, ship
    wb.close()


def chunks(it, size):
    buf = []
    for x in it:
        buf.append(x)
        if len(buf) == size:
            yield buf
            buf = []
    if buf:
        yield buf


def write_file(template_bytes: bytes, items, out: Path) -> None:
    wb = load_workbook(io.BytesIO(template_bytes))
    ws = wb.worksheets[0]
    for r in CLEAR_ROWS:
        for c in CLEAR_COLS:
            ws[f"{c}{r}"].value = None
    for i, (name, url, ship) in enumerate(items):
        r = FIRST_ROW + i
        ws[f"A{r}"].value = name or None
        ws[f"B{r}"].value = url
        ws[f"D{r}"].value = ship or None
    tmp = out.with_name(out.stem + ".tmp.xlsx")
    wb.save(tmp)
    os.replace(tmp, out)             # 다 저장된 뒤에만 진짜 이름으로 바꿈


def main() -> None:
    template, sources = find_files()
    template_bytes = template.read_bytes()
    print(f"양식: {template.name}")
    for src in sources:
        out_dir = OUT_DIR / src.stem
        out_dir.mkdir(parents=True, exist_ok=True)
        for t in out_dir.glob("*.tmp.xlsx"):   # 지난번에 멈추며 남은 임시파일 정리
            t.unlink()
        print(f"\n[{src.name}] 읽는 중... (상품이 많으면 몇 분 걸립니다)")
        made = skipped = total = 0
        for n, items in enumerate(chunks(read_products(src), PER_FILE), 1):
            total += len(items)
            out = out_dir / f"{OUT_PREFIX}{n:04d}.xlsx"
            # 49개가 꽉 찬 파일이 이미 있으면 건너뜀 (마지막 덜 찬 파일은 원본이 늘었을 수 있어 다시 만듦)
            if out.exists() and len(items) == PER_FILE:
                skipped += 1
                continue
            write_file(template_bytes, items, out)
            made += 1
            if made % 100 == 0:
                print(f"  {out.name} 까지 만듦 (상품 {total:,}개 처리)")
        print(f"[{src.name}] 완료: 상품 {total:,}개 → 파일 {made + skipped:,}개 "
              f"(새로 만듦 {made:,}, 이미 있어 건너뜀 {skipped:,})  위치: {out_dir}")
    print("\n모두 끝났습니다.")


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        print("\n중단됨 - 다시 실행하면 이어서 만듭니다.")
