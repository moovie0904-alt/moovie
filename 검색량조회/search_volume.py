# -*- coding: utf-8 -*-
"""
핵심키워드(C열)의 네이버 월간 검색량과 경쟁정도를 네이버 검색광고 API로 조회해서
D~G열에 넣습니다.

준비: 이 폴더의 naver_key.txt 에 3줄을 적어 저장
    CUSTOMER_ID=숫자
    API_KEY=액세스라이선스
    SECRET_KEY=비밀키

사용법:
    python search_volume.py                          (기본 파일)
    python search_volume.py "D:\다른파일.xlsx"
"""
import base64
import hashlib
import hmac
import sys
import time
from pathlib import Path

import requests
from openpyxl import load_workbook

DEFAULT_INPUT = r"C:\자동화시스템\합친결과_키워드.xlsx"
KEYWORD_COL = 3                      # C열 = 핵심키워드
OUT_HEADERS = ["PC검색량", "모바일검색량", "총검색량", "경쟁정도"]   # D~G열
BASE_URL = "https://api.searchad.naver.com"
URI = "/keywordstool"
HERE = Path(__file__).resolve().parent


def load_keys() -> dict:
    path = HERE / "naver_key.txt"
    if not path.exists():
        sys.exit(f"키 파일이 없습니다: {path}\n메모장으로 만들어 CUSTOMER_ID=, API_KEY=, SECRET_KEY= 3줄을 적어주세요.")
    keys = {}
    for line in path.read_text(encoding="utf-8-sig").splitlines():
        if "=" in line:
            k, v = line.split("=", 1)
            keys[k.strip().upper()] = v.strip()
    for need in ("CUSTOMER_ID", "API_KEY", "SECRET_KEY"):
        if not keys.get(need):
            sys.exit(f"naver_key.txt 에 {need}= 값이 없습니다.")
    return keys


def headers(keys: dict) -> dict:
    ts = str(int(time.time() * 1000))
    msg = f"{ts}.GET.{URI}".encode()
    sig = base64.b64encode(hmac.new(keys["SECRET_KEY"].encode(), msg, hashlib.sha256).digest()).decode()
    return {"X-Timestamp": ts, "X-API-KEY": keys["API_KEY"], "X-Customer": keys["CUSTOMER_ID"], "X-Signature": sig}


def to_int(v) -> int:
    """'< 10' 처럼 오는 값은 5로 봄."""
    if isinstance(v, (int, float)):
        return int(v)
    return 5 if "<" in str(v) else int(str(v).replace(",", "") or 0)


def normalize(k: str) -> str:
    return "".join(str(k).split()).upper()


def fetch(keys: dict, keywords: list[str]) -> dict[str, tuple]:
    """키워드 최대 5개를 한 번에 조회. {정규화키워드: (pc, mobile, 경쟁정도)}"""
    for attempt in range(3):
        r = requests.get(BASE_URL + URI, headers=headers(keys),
                         params={"hintKeywords": ",".join(keywords), "showDetail": "1"}, timeout=20)
        if r.status_code == 429:          # 너무 빠름 → 잠깐 쉬고 다시
            time.sleep(2 * (attempt + 1))
            continue
        if r.status_code in (401, 403):
            sys.exit(f"인증 실패 ({r.status_code}): naver_key.txt 의 키 3개를 다시 확인하세요.\n{r.text[:200]}")
        r.raise_for_status()
        out = {}
        for item in r.json().get("keywordList", []):
            out[normalize(item["relKeyword"])] = (
                to_int(item.get("monthlyPcQcCnt", 0)),
                to_int(item.get("monthlyMobileQcCnt", 0)),
                item.get("compIdx", ""),
            )
        return out
    raise RuntimeError("요청이 계속 거절됩니다. 잠시 후 다시 실행하세요.")


def main() -> None:
    in_path = Path(sys.argv[1] if len(sys.argv) > 1 else DEFAULT_INPUT)
    if not in_path.exists():
        sys.exit(f"파일을 찾을 수 없습니다: {in_path}")
    keys = load_keys()

    print(f"읽는 중: {in_path} (상품이 많으면 1~2분 걸립니다)")
    wb = load_workbook(in_path)
    ws = wb.active
    for i, h in enumerate(OUT_HEADERS):
        ws.cell(1, KEYWORD_COL + 1 + i).value = h

    # 키워드가 있는 줄만 모음 (같은 키워드는 한 번만 조회)
    rows = {}
    for r in range(2, ws.max_row + 1):
        kw = ws.cell(r, KEYWORD_COL).value
        if kw and str(kw).strip():
            rows.setdefault(str(kw).strip(), []).append(r)
    if not rows:
        sys.exit("C열에 핵심키워드가 없습니다.")
    unique = list(rows)
    print(f"조회할 키워드 {len(unique)}개")

    results = {}
    for start in range(0, len(unique), 5):
        chunk = unique[start:start + 5]
        # 네이버 API는 키워드에 띄어쓰기가 있으면 거절하므로 붙여서 보냄
        got = fetch(keys, [k.replace(" ", "") for k in chunk])
        for k in chunk:
            results[k] = got.get(normalize(k))
        print(f"  {min(start + 5, len(unique))}/{len(unique)}")
        time.sleep(0.3)

    for kw, rs in rows.items():
        res = results.get(kw)
        for r in rs:
            if res:
                pc, mo, comp = res
                vals = [pc, mo, pc + mo, comp]
            else:
                vals = ["", "", "", "조회안됨"]
            for i, v in enumerate(vals):
                ws.cell(r, KEYWORD_COL + 1 + i).value = v

    out = in_path.with_name(in_path.stem + "_검색량.xlsx")
    wb.save(out)
    print(f"\n저장 완료: {out}")
    print("※ 검색량이 10 미만인 키워드는 5로 표시됩니다.")


if __name__ == "__main__":
    main()
